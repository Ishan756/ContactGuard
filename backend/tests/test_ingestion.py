import io

import docx
import fitz
import pytest
from httpx import AsyncClient


def create_synthetic_pdf() -> bytes:
    """Generate a clean synthetic PDF contract with 2 pages and standard section headings."""
    doc = fitz.open()

    # Page 1
    page1 = doc.new_page(width=595, height=842)
    p1_text = (
        "MASTER SERVICES AGREEMENT\n\n"
        "ARTICLE 1. CONFIDENTIALITY\n"
        "Each party agrees to hold the Confidential Information of the other in strict confidence "
        "and not to disclose such Confidential Information to any third party without prior consent.\n\n"
        "ARTICLE 2. LIMITATION OF LIABILITY\n"
        "In no event shall either party's aggregate liability exceed the total fees paid under this Agreement."
    )
    page1.insert_text((50, 72), p1_text, fontsize=11)

    # Page 2
    page2 = doc.new_page(width=595, height=842)
    p2_text = (
        "ARTICLE 3. TERM AND TERMINATION\n"
        "This Agreement commences on the Effective Date and continues for twelve (12) months.\n\n"
        "ARTICLE 4. GOVERNING LAW\n"
        "This Agreement shall be construed and governed in accordance with the laws of Delaware."
    )
    page2.insert_text((50, 72), p2_text, fontsize=11)

    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def create_scanned_synthetic_pdf() -> bytes:
    """Generate a synthetic PDF with an image/shape but no extractable text characters (< 50 chars)."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    # Draw a rectangle mimicking a scanned page container
    page.draw_rect(fitz.Rect(50, 50, 545, 792), color=(0.2, 0.2, 0.2), fill=(0.95, 0.95, 0.95))
    page.insert_text((60, 60), "P", fontsize=8)  # Single character
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def create_encrypted_pdf() -> bytes:
    """Generate a password-encrypted PDF."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 72), "Confidential Protected Agreement", fontsize=11)
    pdf_bytes = doc.tobytes(
        encryption=fitz.PDF_ENCRYPT_AES_256,
        user_pw="secure_password_123",
        owner_pw="owner_password_123",
    )
    doc.close()
    return pdf_bytes


def create_synthetic_docx() -> bytes:
    """Generate a clean synthetic DOCX contract with headings and paragraphs."""
    doc = docx.Document()
    doc.add_heading("NON-DISCLOSURE AGREEMENT", level=0)

    doc.add_heading("SECTION 1. DEFINITIONS", level=1)
    doc.add_paragraph(
        "Confidential Information shall mean all proprietary, technical, and commercial data "
        "disclosed by the Disclosing Party to the Receiving Party."
    )

    doc.add_heading("SECTION 2. NON-USE OBLIGATIONS", level=1)
    doc.add_paragraph(
        "The Receiving Party shall not use any Confidential Information for any purpose other than "
        "evaluating the potential business partnership."
    )

    file_stream = io.BytesIO()
    doc.save(file_stream)
    return file_stream.getvalue()


@pytest.mark.asyncio
async def test_upload_and_extract_valid_pdf(async_client: AsyncClient):
    """Test successful PDF upload, metadata extraction, and chunk persistence."""
    pdf_bytes = create_synthetic_pdf()
    files = {"file": ("enterprise_msa.pdf", pdf_bytes, "application/pdf")}

    response = await async_client.post("/api/v1/documents", files=files)
    assert response.status_code == 201
    data = response.json()

    assert data["document"]["original_filename"] == "enterprise_msa.pdf"
    assert data["document"]["status"] == "completed"
    assert data["document"]["total_pages"] == 2
    assert data["document"]["total_characters"] > 200
    assert data["chunks_extracted"] >= 1
    doc_id = data["document"]["id"]

    # Verify retrieval by ID
    get_resp = await async_client.get(f"/api/v1/documents/{doc_id}")
    assert get_resp.status_code == 200
    get_data = get_resp.json()
    assert get_data["id"] == doc_id
    assert get_data["file_hash_sha256"] is not None

    # Verify text chunks with pagination
    chunks_resp = await async_client.get(f"/api/v1/documents/{doc_id}/chunks?page=1&page_size=10")
    assert chunks_resp.status_code == 200
    chunks_data = chunks_resp.json()
    assert chunks_data["total_chunks"] >= 1
    first_chunk = chunks_data["chunks"][0]
    assert first_chunk["document_id"] == doc_id
    assert first_chunk["page_number"] in (1, 2)
    assert first_chunk["source_identifier"].startswith(f"doc:{doc_id}:p")


@pytest.mark.asyncio
async def test_upload_and_extract_valid_docx(async_client: AsyncClient):
    """Test successful DOCX upload and section extraction."""
    docx_bytes = create_synthetic_docx()
    files = {
        "file": (
            "bilateral_nda.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }

    response = await async_client.post("/api/v1/documents", files=files)
    assert response.status_code == 201
    data = response.json()

    assert data["document"]["original_filename"] == "bilateral_nda.docx"
    assert data["document"]["status"] == "completed"
    assert data["chunks_extracted"] >= 2
    doc_id = data["document"]["id"]

    # Verify chunks contain detected section headings
    chunks_resp = await async_client.get(f"/api/v1/documents/{doc_id}/chunks")
    assert chunks_resp.status_code == 200
    chunks = chunks_resp.json()["chunks"]
    headings = [c["section_heading"] for c in chunks if c["section_heading"]]
    assert len(headings) >= 1


@pytest.mark.asyncio
async def test_reject_empty_document(async_client: AsyncClient):
    """Test rejection of 0-byte file upload."""
    files = {"file": ("empty.pdf", b"", "application/pdf")}
    response = await async_client.post("/api/v1/documents", files=files)
    assert response.status_code == 422
    data = response.json()
    assert data["error"]["code"] == "VALIDATION_ERROR"
    assert "empty" in data["error"]["message"].lower()


@pytest.mark.asyncio
async def test_reject_unsupported_file_extension(async_client: AsyncClient):
    """Test rejection of non-allowed file extension."""
    files = {"file": ("script.exe", b"binarycontent", "application/octet-stream")}
    response = await async_client.post("/api/v1/documents", files=files)
    assert response.status_code == 422
    data = response.json()
    assert data["error"]["code"] == "VALIDATION_ERROR"
    assert "unsupported" in data["error"]["message"].lower()


@pytest.mark.asyncio
async def test_reject_spoofed_content_mismatch(async_client: AsyncClient):
    """Test rejection of file named .pdf containing plain non-PDF text."""
    files = {"file": ("spoofed.pdf", b"this is definitely not a real pdf", "application/pdf")}
    response = await async_client.post("/api/v1/documents", files=files)
    assert response.status_code == 422
    data = response.json()
    assert "signature" in data["error"]["message"].lower()


@pytest.mark.asyncio
async def test_detect_scanned_pdf_needs_ocr(async_client: AsyncClient):
    """Test that image-only / scanned PDFs are flagged with needs_ocr status."""
    scanned_pdf = create_scanned_synthetic_pdf()
    files = {"file": ("scanned_agreement.pdf", scanned_pdf, "application/pdf")}

    response = await async_client.post("/api/v1/documents", files=files)
    assert response.status_code == 201
    data = response.json()

    assert data["document"]["status"] == "needs_ocr"
    assert "ocr is required" in data["document"]["error_message"].lower()


@pytest.mark.asyncio
async def test_handle_encrypted_pdf_gracefully(async_client: AsyncClient):
    """Test encrypted PDF handling reports password requirement without crashing."""
    encrypted_pdf = create_encrypted_pdf()
    files = {"file": ("protected_agreement.pdf", encrypted_pdf, "application/pdf")}

    response = await async_client.post("/api/v1/documents", files=files)
    assert response.status_code == 201
    data = response.json()

    assert data["document"]["status"] == "failed"
    assert "encrypted" in data["document"]["error_message"].lower()


@pytest.mark.asyncio
async def test_list_documents_pagination(async_client: AsyncClient):
    """Test listing documents returns count and metadata items."""
    response = await async_client.get("/api/v1/documents?limit=10&offset=0")
    assert response.status_code == 200
    data = response.json()
    assert "documents" in data
    assert "total" in data
    assert isinstance(data["documents"], list)
